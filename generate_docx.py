import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_document():
    doc = docx.Document()

    # Set page margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Styles
    COLOR_PRIMARY = RGBColor(15, 23, 42)      # Slate 900
    COLOR_ACCENT = RGBColor(30, 58, 138)      # Blue 900
    COLOR_SECONDARY = RGBColor(71, 85, 105)   # Slate 600
    COLOR_BODY = RGBColor(30, 41, 59)         # Slate 800

    # Document Header / Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("MODUL MATERI LENGKAP: BIG DATA (P7 - P9)")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_ACCENT

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(16)
    run_sub = sub_p.add_run("HDFS Storage Architecture, YARN Resource Management, & MapReduce Processing Model")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(11)
    run_sub.font.italic = True
    run_sub.font.color.rgb = COLOR_SECONDARY

    # Divider line
    div_p = doc.add_paragraph()
    div_p.paragraph_format.space_after = Pt(18)
    pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="18" w:space="1" w:color="1E3A8A"/></w:pBdr>')
    div_p._p.get_or_add_pPr().append(pBdr)

    # Callout Box: Overview
    overview_table = doc.add_table(rows=1, cols=1)
    overview_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = overview_table.cell(0, 0)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=140, bottom=140, left=180, right=180)
    p_box = cell.paragraphs[0]
    p_box.paragraph_format.space_after = Pt(0)
    r_box_title = p_box.add_run("Ringkasan Dokumen:\n")
    r_box_title.bold = True
    r_box_title.font.size = Pt(10.5)
    r_box_title.font.color.rgb = COLOR_PRIMARY
    r_box_body = p_box.add_run(
        "Dokumen ini dirancang sebagai catatan materi terstruktur dan komprehensif untuk persiapan ujian "
        "mata kuliah Big Data materi Pertemuan 7 sampai Pertemuan 9. Dokumen mencakup teori arsitektur, "
        "rumus matematis, alokasi blok disk HDFS, siklus eksekusi MapReduce, dan rangkuman pembahasan bank soal."
    )
    r_box_body.font.size = Pt(9.5)
    r_box_body.font.color.rgb = COLOR_SECONDARY

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Function to add styled Section Heading
    def add_heading_1(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(8)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = COLOR_ACCENT
        # Add bottom border
        pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="2" w:color="CBD5E1"/></w:pBdr>')
        h._p.get_or_add_pPr().append(pBdr)
        return h

    def add_heading_2(text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(12)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
        run = h.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = COLOR_PRIMARY
        return h

    def add_body(text, bold_prefix=None, space_after=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = "Arial"
            r_bold.font.size = Pt(10)
            r_bold.font.bold = True
            r_bold.font.color.rgb = COLOR_PRIMARY
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = COLOR_BODY
        return p

    def add_bullet(text, bold_prefix=None):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        if bold_prefix:
            r_bold = p.add_run(bold_prefix)
            r_bold.font.name = "Arial"
            r_bold.font.size = Pt(10)
            r_bold.font.bold = True
            r_bold.font.color.rgb = COLOR_PRIMARY
        r = p.add_run(text)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.color.rgb = COLOR_BODY
        return p

    # --- MODUL 1 ---
    add_heading_1("MODUL 1: FONDASI BIG DATA & EVOLUSI EKOSISTEM HADOOP")
    add_body("Big Data muncul sebagai respons terhadap ketidakmampuan basis data relasional (RDBMS) tradisional dalam menangani data skala petabyte dengan pertumbuhan eksponensial.")
    add_heading_2("1.1 Karakteristik 5V Big Data")
    add_bullet("Skala ukuran data yang masif (terabyte hingga petabyte/exabyte).", bold_prefix="Volume: ")
    add_bullet("Kecepatan aliran data masuk dan kebutuhan pemrosesan (batch vs real-time).", bold_prefix="Velocity: ")
    add_bullet("Keragaman format data: terstruktur (SQL), semi-terstruktur (JSON, XML), dan tidak terstruktur (video, teks, log).", bold_prefix="Variety: ")
    add_bullet("Tingkat kualitas, akurasi, dan integritas data yang dapat dipercaya.", bold_prefix="Veracity: ")
    add_bullet("Nilai bisnis yang diekstraksi dari hasil analisis data.", bold_prefix="Value: ")

    add_heading_2("1.2 Paradigma Skalabilitas: Scale-Up vs Scale-Out")
    add_body(
        "Pendekatan tradisional (Scale-Up / Vertical) membesarkan spesifikasi satu mesin tunggal (menambah RAM, CPU, disk). "
        "Kelemahannya adalah batas fisik hardware dan biaya yang sangat mahal. Sebaliknya, Apache Hadoop mengadopsi Scale-Out (Horizontal Scaling), "
        "yaitu menggabungkan ratusan hingga ribuan komputer komoditas (commodity hardware) berbiaya terjangkau menjadi satu klaster terpadu."
    )

    add_heading_2("1.3 Sejarah & Google Whitepapers")
    add_body(
        "Hadoop diciptakan oleh Doug Cutting dan Mike Cafarella, dinamai dari mainan gajah anak Doug Cutting. "
        "Hadoop didasari oleh dua makalah teknis fundamental dari Google Research:\n"
        "1. Google File System (GFS) (2003) -> Menginspirasi lahirnya HDFS.\n"
        "2. Google MapReduce (2004) -> Menginspirasi paradigma komputasi MapReduce."
    )

    add_heading_2("1.4 Perbandingan Hadoop 1.x vs Hadoop 2.x/3.x")
    add_bullet("Hadoop 1.x menggabungkan manajemen klaster dan komputasi di JobTracker. Menyebabkan Single Point of Failure (SPOF), keterbatasan skalabilitas (~4.000 node), dan hanya bisa menjalankan MapReduce.", bold_prefix="Hadoop 1.x: ")
    add_bullet("Hadoop 2.x memisahkan komputasi dari resource management dengan memperkenalkan YARN. Klaster dapat menjalankan Spark, Flink, Tez, dan HBase secara berdampingan.", bold_prefix="Hadoop 2.x: ")
    add_bullet("Hadoop 3.x menambahkan Erasure Coding (menghemat storage hingga 50%), multi-standby NameNodes, dan dukungan GPU Container.", bold_prefix="Hadoop 3.x: ")

    # --- MODUL 2 ---
    add_heading_1("MODUL 2: HDFS (HADOOP DISTRIBUTED FILE SYSTEM) DEEP-DIVE")
    add_body(
        "HDFS adalah sistem berkas terdistribusi berbasis blok dengan karakteristik Write-Once-Read-Many (WORM), "
        "kecepatan throughput streaming data yang tinggi, dan toleransi kegagalan otomatis."
    )

    add_heading_2("2.1 Konsep HDFS Block & Alokasi Fisik Cakram (CRITICAL CONCEPT)")
    add_body(
        "Data di HDFS dipecah menjadi blok-blok besar (default 128 MB di Hadoop 2/3 atau 64 MB di Hadoop 1). "
        "Ukuran blok besar bertujuan meminimalkan disk seek time dan menjaga efisiensi RAM NameNode."
    )

    # Highlight box for the 514 MB question
    box_calc = doc.add_table(rows=1, cols=1)
    box_calc.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_calc = box_calc.cell(0, 0)
    set_cell_background(c_calc, "ECFDF5")
    set_cell_margins(c_calc, top=120, bottom=120, left=160, right=160)
    p_c = c_calc.paragraphs[0]
    p_c.paragraph_format.space_after = Pt(0)
    r_t = p_c.add_run("Studi Kasus Kuis (Gambar): Alokasi Fisik Blok Terakhir\n")
    r_t.bold = True
    r_t.font.color.rgb = RGBColor(6, 95, 70)
    r_b = p_c.add_run(
        "Soal: Berapa ukuran ruang cakram fisik yang terpakai di HDFS untuk blok terakhir dari berkas berukuran 514 MB dengan ukuran blok 128 MB?\n\n"
        "Analisis Perhitungan:\n"
        "• Ukuran berkas = 514 MB, ukuran blok = 128 MB.\n"
        "• 4 blok pertama = 4 x 128 MB = 512 MB (terisi penuh).\n"
        "• Blok ke-5 (terakhir) = 514 MB - 512 MB = 2 MB sisa data.\n"
        "• Di HDFS, blok adalah batas alokasi logis maksimum, BUKAN alokasi fisik statis di disk.\n"
        "-> Jawaban Benar: Ruang cakram fisik yang terpakai adalah 2 MB saja (bukan 128 MB penuh)!"
    )
    r_b.font.size = Pt(9.5)
    r_b.font.color.rgb = RGBColor(4, 120, 87)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    add_heading_2("2.2 Komponen Arsitektur HDFS")
    add_bullet("Master daemon. Mengelola namespace sistem berkas dan struktur pohon direktori. Menyimpan seluruh metadata di RAM untuk akses instan. Memiliki dua berkas persisten di disk: fsimage (snapshot) dan edits/editlog (catatan transaksi).", bold_prefix="1. NameNode: ")
    add_bullet("Worker daemon. Menyimpan blok data fisik aktual di disk lokal host. Mengirimkan Heartbeat setiap 3 detik dan Block Report setiap 6 jam ke NameNode.", bold_prefix="2. DataNode: ")
    add_bullet("BUKAN hot standby/failover! Bertugas murni melakukan Checkpointing secara periodik (menggabungkan fsimage dan editlog agar berkas log tidak membengkak).", bold_prefix="3. Secondary NameNode: ")
    add_bullet("Menghilangkan SPOF NameNode dengan sepasang Active NameNode dan Standby NameNode yang disinkronisasi melalui Quorum Journal Manager (QJM) dan diawasi oleh ZooKeeper Failover Controller (ZKFC).", bold_prefix="4. High Availability (HA): ")

    add_heading_2("2.3 Strategi Replikasi & Rack Awareness (Default Replication Factor = 3)")
    add_bullet("Replika 1 ditaruh di node lokal (tempat klien menulis data jika di dalam klaster).", bold_prefix="Aturan Penempatan 1: ")
    add_bullet("Replika 2 ditaruh di node pada rak yang berbeda (off-rack).", bold_prefix="Aturan Penempatan 2: ")
    add_bullet("Replika 3 ditaruh di node yang berbeda pada rak yang sama dengan replika ke-2.", bold_prefix="Aturan Penempatan 3: ")
    add_body("Strategi ini mengamankan data jika terjadi kegagalan satu rak server penuh, sekaligus mengoptimalkan bandwidth switch antar-rak.")

    # --- MODUL 3 ---
    add_heading_1("MODUL 3: YARN (YET ANOTHER RESOURCE NEGOTIATOR)")
    add_body("YARN bertindak sebagai sistem operasi terdistribusi klaster yang mengelola alokasi sumber daya (CPU core & RAM) secara efisien.")
    add_bullet("Otoritas tertinggi manajemen resource. Terdiri dari Scheduler (alokasi murni kapasitas tanpa memantau status aplikasi: FIFO, Capacity, Fair Scheduler) dan ApplicationsManager/ASM (menerima pengajuan aplikasi dan meluncurkan ApplicationMaster).", bold_prefix="1. ResourceManager (RM): ")
    add_bullet("Daemon di setiap worker node yang memantau penggunaan resource lokal dan mengelola siklus Container.", bold_prefix="2. NodeManager (NM): ")
    add_bullet("Entitas unik per aplikasi yang menegosiasikan container ke RM dan bekerja sama dengan NM untuk eksekusi tugas.", bold_prefix="3. ApplicationMaster (AM): ")
    add_bullet("Abstraksi alokasi sumber daya fisik (misal: 2 GB RAM + 1 vCore) tempat task dijalankan.", bold_prefix="4. Container: ")

    # --- MODUL 4 ---
    add_heading_1("MODUL 4: PARADIGMA & ALUR KERJA MAPREDUCE")
    add_heading_2("4.1 Filosofi Data Locality")
    add_body(
        "Prinsip utama: Moving computation is cheaper than moving data (memindahkan komputasi yang ringan lebih murah "
        "daripada memindahkan data ratusan megabyte melalui jaringan klaster).\n"
        "Tingkatan efisiensi: 1. Node-local (Data-local) > 2. Rack-local > 3. Off-rack (Different rack)."
    )

    add_heading_2("4.2 Tahapan Alur Eksekusi MapReduce")
    add_bullet("Pembagian berkas input secara logis (umumnya 1 split = 1 blok HDFS = 1 Mapper task).", bold_prefix="1. InputSplit: ")
    add_bullet("Membaca byte mentah dari split menjadi pasangan key-value pertama <K1, V1> (misal: byte offset dan teks baris).", bold_prefix="2. RecordReader: ")
    add_bullet("Mengeksekusi logika pemrosesan baris dan menghasilkan intermediate key-value <K2, V2>.", bold_prefix="3. Map Phase: ")
    add_bullet("Mini-Reducer lokal di node mapper untuk melakukan pra-agregasi lokal (mengurangi trafik jaringan saat shuffle). Hanya aman untuk operasi asosiatif & komutatif seperti Sum dan Count.", bold_prefix="4. Combiner (Opsional): ")
    add_bullet("Menentukan Reducer tujuan untuk setiap key (default: (key.hashCode() & Integer.MAX_VALUE) % numReduceTasks). Menjamin key yang sama selalu masuk ke Reducer yang sama.", bold_prefix="5. Partitioner: ")
    add_bullet("Proses mengunduh partisi data dari semua Mapper ke Reducer via jaringan (Shuffle) dan mengurutkan key secara otomatis (Sort/Merge).", bold_prefix="6. Shuffle and Sort: ")
    add_bullet("Mengagregasi list of values untuk setiap key unik <K2, List<V2>> menghasilkan output akhir <K3, V3>.", bold_prefix="7. Reduce Phase: ")
    add_bullet("Menulis output akhir ke berkas part-r-00000 di HDFS.", bold_prefix="8. RecordWriter: ")

    # --- MODUL 5 ---
    add_heading_1("MODUL 5: FORMULA RUMUS & STUDI KASUS UJIAN")
    
    # Table of formulas
    tbl = doc.add_table(rows=5, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_headers = ["Parameter Perhitungan", "Formula / Rumus", "Contoh Kasus Ujian"]
    for j, h_text in enumerate(tbl_headers):
        c = tbl.cell(0, j)
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, top=100, bottom=100, left=120, right=120)
        p = c.paragraphs[0]
        r = p.add_run(h_text)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    data_rows = [
        ("Alokasi Fisik Blok Terakhir", "Sisa = Ukuran Berkas mod Ukuran Blok", "File 514 MB, blok 128 MB -> 2 MB"),
        ("Jumlah Blok Dibutuhkan", "N = ceil(Ukuran Berkas / Ukuran Blok)", "File 300 MB, blok 128 MB -> 3 blok"),
        ("Total Storage dengan Replikasi", "Total = Ukuran Berkas x Faktor Replikasi", "File 100 GB, RF = 3 -> 300 GB"),
        ("Jumlah Mapper Default", "Jumlah Mapper = Total InputSplit (Blok)", "File 1024 MB, blok 128 MB -> 8 Mapper")
    ]

    for i, (p_name, f_name, c_case) in enumerate(data_rows, start=1):
        bg = "F8FAFC" if i % 2 == 1 else "FFFFFF"
        for j, text in enumerate([p_name, f_name, c_case]):
            c = tbl.cell(i, j)
            set_cell_background(c, bg)
            set_cell_margins(c, top=80, bottom=80, left=100, right=100)
            p = c.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9)
            r.font.color.rgb = COLOR_BODY

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # --- MODUL 6 ---
    add_heading_1("MODUL 6: KUNCI JAWABAN & PEMBAHASAN 50 SOAL KUIS")
    add_body("Berikut adalah ringkasan kunci jawaban dan konsep esensial dari 50 soal kuis persiapan UTS:")

    # Import questions from soal_bank.json
    import json
    with open("c:/Users/Satria Prana/Documents/Programing/Big Data/belajar big data/soal_bank.json", "r", encoding="utf-8") as f:
        questions = json.load(f)

    for q in questions:
        p_q = doc.add_paragraph()
        p_q.paragraph_format.space_before = Pt(8)
        p_q.paragraph_format.space_after = Pt(2)
        p_q.paragraph_format.keep_with_next = True
        r_num = p_q.add_run(f"Soal #{q['id']}: ")
        r_num.bold = True
        r_num.font.color.rgb = COLOR_ACCENT
        p_q.add_run(q['question']).font.bold = True

        correct_opt = q['options'][q['correctAnswer']]
        p_ans = doc.add_paragraph()
        p_ans.paragraph_format.space_after = Pt(2)
        r_k = p_ans.add_run("Kunci Jawaban: ")
        r_k.bold = True
        r_k.font.color.rgb = RGBColor(16, 185, 129)
        p_ans.add_run(correct_opt).font.bold = True

        p_exp = doc.add_paragraph()
        p_exp.paragraph_format.space_after = Pt(6)
        r_e = p_exp.add_run("Pembahasan: ")
        r_e.italic = True
        r_e.font.color.rgb = COLOR_SECONDARY
        r_desc = p_exp.add_run(q['explanation'])
        r_desc.font.size = Pt(9.5)
        r_desc.font.color.rgb = COLOR_BODY

    # Footer note
    footer = doc.sections[0].footer
    f_p = footer.paragraphs[0]
    f_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    f_r = f_p.add_run("Dokumen Materi & Bank Soal Big Data P7-P9 | Antigravity AI")
    f_r.font.size = Pt(8.5)
    f_r.font.italic = True
    f_r.font.color.rgb = RGBColor(148, 163, 184)

    output_path = "c:/Users/Satria Prana/Documents/Programing/Big Data/belajar big data/DOKUMEN_MATERI_LENGKAP_BIG_DATA_P7_P9.docx"
    try:
        doc.save(output_path)
        print("Document successfully created at:", output_path)
    except PermissionError:
        fallback_path = "c:/Users/Satria Prana/Documents/Programing/Big Data/belajar big data/DOKUMEN_MATERI_LENGKAP_BIG_DATA_P7_P9_BANK250.docx"
        doc.save(fallback_path)
        print("Original docx is currently open in Word. Saved to:", fallback_path)

if __name__ == "__main__":
    create_document()
